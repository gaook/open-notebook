"""文档解析器 — 从 PDF/DOCX/TXT/MD 提取纯文本"""

from pathlib import Path


def parse_document(file_path: str, file_type: str) -> str:
    """提取文档纯文本内容"""
    if file_type == "pdf":
        return _parse_pdf(file_path)
    elif file_type == "docx":
        return _parse_docx(file_path)
    elif file_type in ("txt", "md"):
        return Path(file_path).read_text(encoding="utf-8")
    else:
        raise ValueError(f"不支持的文件类型: {file_type}")


def _parse_pdf(file_path: str) -> str:
    import fitz  # PyMuPDF

    doc = fitz.open(file_path)
    texts = []
    for page in doc:
        texts.append(page.get_text("text"))
    doc.close()
    return "\n".join(texts)


def _parse_docx(file_path: str) -> str:
    from docx import Document

    doc = Document(file_path)
    texts = []
    for para in doc.paragraphs:
        if para.text.strip():
            texts.append(para.text)
    # 也提取表格内容
    for table in doc.tables:
        for row in table.rows:
            row_text = "\t".join(cell.text for cell in row.cells)
            if row_text.strip():
                texts.append(row_text)
    return "\n".join(texts)
